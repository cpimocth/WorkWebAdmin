from __future__ import annotations

import json
import os
import re
import threading
import time
import webbrowser
import zipfile
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
import xml.etree.ElementTree as ET

BASE_DIR = Path(__file__).resolve().parent
EXCEL_PATH = BASE_DIR / 'Database_CPI.xlsx'
HOST = '127.0.0.1'
PORT = 8765


def _clean(v):
    if v is None:
        return ''
    return v


def _read_with_openpyxl(path: Path):
    try:
        from openpyxl import load_workbook
    except Exception:
        return None
    wb = load_workbook(path, read_only=True, data_only=True)
    def sheet_rows(name):
        if name not in wb.sheetnames:
            return []
        ws = wb[name]
        rows = ws.iter_rows(values_only=True)
        try:
            headers = [str(v).strip() if v is not None else '' for v in next(rows)]
        except StopIteration:
            return []
        out = []
        for vals in rows:
            if all(v is None for v in vals):
                continue
            row = {}
            for h, v in zip(headers, vals):
                if not h:
                    continue
                if hasattr(v, 'isoformat'):
                    try:
                        v = v.isoformat()
                    except Exception:
                        pass
                row[h] = _clean(v)
            out.append(row)
        return out
    return {'users': sheet_rows('Users'), 'tasks': sheet_rows('Tasks')}


def _col_index(cell_ref: str) -> int:
    m = re.match(r'([A-Z]+)', cell_ref or '')
    if not m:
        return 0
    n = 0
    for ch in m.group(1):
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _xml_text(node):
    if node is None:
        return ''
    return ''.join(node.itertext())


def _read_xlsx_stdlib(path: Path):
    NS_MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    NS_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    NS_PKG_REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
    with zipfile.ZipFile(path) as z:
        shared = []
        if 'xl/sharedStrings.xml' in z.namelist():
            root = ET.fromstring(z.read('xl/sharedStrings.xml'))
            for si in root.findall(f'{{{NS_MAIN}}}si'):
                shared.append(_xml_text(si))

        wb_root = ET.fromstring(z.read('xl/workbook.xml'))
        rel_root = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        rels = {}
        for rel in rel_root.findall(f'{{{NS_PKG_REL}}}Relationship'):
            rels[rel.attrib.get('Id')] = rel.attrib.get('Target', '')

        sheet_targets = {}
        sheets = wb_root.find(f'{{{NS_MAIN}}}sheets')
        if sheets is not None:
            for sh in sheets.findall(f'{{{NS_MAIN}}}sheet'):
                name = sh.attrib.get('name', '')
                rid = sh.attrib.get(f'{{{NS_REL}}}id')
                target = rels.get(rid, '')
                if target.startswith('/'):
                    full = target.lstrip('/')
                else:
                    full = 'xl/' + target.lstrip('/')
                full = os.path.normpath(full).replace('\\', '/')
                sheet_targets[name] = full

        def cell_value(c):
            t = c.attrib.get('t', '')
            if t == 'inlineStr':
                is_node = c.find(f'{{{NS_MAIN}}}is')
                return _xml_text(is_node)
            v = c.find(f'{{{NS_MAIN}}}v')
            txt = '' if v is None or v.text is None else v.text
            if t == 's':
                try:
                    return shared[int(txt)]
                except Exception:
                    return ''
            if t == 'b':
                return 'TRUE' if txt == '1' else 'FALSE'
            if t in ('str', 'e'):
                return txt
            # Numeric cells: preserve integer-looking values without .0
            if txt == '':
                return ''
            try:
                f = float(txt)
                if f.is_integer():
                    return str(int(f))
                return txt
            except Exception:
                return txt

        def sheet_rows(name):
            target = sheet_targets.get(name)
            if not target or target not in z.namelist():
                return []
            root = ET.fromstring(z.read(target))
            sheet_data = root.find(f'{{{NS_MAIN}}}sheetData')
            if sheet_data is None:
                return []
            matrix = []
            for row_node in sheet_data.findall(f'{{{NS_MAIN}}}row'):
                vals = {}
                max_idx = -1
                for c in row_node.findall(f'{{{NS_MAIN}}}c'):
                    idx = _col_index(c.attrib.get('r', ''))
                    vals[idx] = cell_value(c)
                    max_idx = max(max_idx, idx)
                if max_idx < 0:
                    matrix.append([])
                else:
                    matrix.append([vals.get(i, '') for i in range(max_idx + 1)])
            if not matrix:
                return []
            headers = [str(v).strip() for v in matrix[0]]
            out = []
            for vals in matrix[1:]:
                if not any(str(v).strip() for v in vals):
                    continue
                row = {}
                for i, h in enumerate(headers):
                    if not h:
                        continue
                    row[h] = vals[i] if i < len(vals) else ''
                out.append(row)
            return out

        return {'users': sheet_rows('Users'), 'tasks': sheet_rows('Tasks')}


def read_database():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f'ไม่พบ {EXCEL_PATH.name} ในโฟลเดอร์เดียวกับ server.py')
    data = _read_with_openpyxl(EXCEL_PATH)
    if data is None:
        data = _read_xlsx_stdlib(EXCEL_PATH)
    if not data.get('users'):
        raise ValueError('ไม่พบข้อมูล Users ใน Database_CPI.xlsx')
    return data


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(BASE_DIR), **kwargs)

    def end_headers(self):
        self.send_header('Cache-Control', 'no-store, no-cache, must-revalidate, max-age=0')
        self.send_header('Pragma', 'no-cache')
        self.send_header('Expires', '0')
        super().end_headers()

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == '/api/db':
            try:
                data = read_database()
                body = json.dumps(data, ensure_ascii=False).encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            except Exception as e:
                body = json.dumps({'error': str(e)}, ensure_ascii=False).encode('utf-8')
                self.send_response(500)
                self.send_header('Content-Type', 'application/json; charset=utf-8')
                self.send_header('Content-Length', str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            return
        return super().do_GET()

    def log_message(self, fmt, *args):
        print('[CPI]', fmt % args)


def open_browser():
    time.sleep(0.8)
    webbrowser.open(f'http://{HOST}:{PORT}/index.html')


if __name__ == '__main__':
    print('=' * 64)
    print(' TPSO CPI Local Web Server')
    print(f' Excel : {EXCEL_PATH}')
    print(f' URL   : http://{HOST}:{PORT}/index.html')
    print(' เมื่อแก้ Database_CPI.xlsx ให้ Save แล้ว Refresh หน้าเว็บ')
    print(' ปิดระบบโดยปิดหน้าต่างนี้ หรือกด Ctrl+C')
    print('=' * 64)
    threading.Thread(target=open_browser, daemon=True).start()
    try:
        ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
    except KeyboardInterrupt:
        pass
